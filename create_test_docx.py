#!/usr/bin/env python3
"""
Script to create a test Word document for extraction testing
"""
import sys
import os

try:
    from docx import Document
    from docx.shared import Inches
    
    def create_test_document():
        """Create a test Word document with contract content"""
        doc = Document()
        
        title = doc.add_heading('CONTRACT AGREEMENT', 0)
        
        doc.add_paragraph('Date: January 15, 2024')
        doc.add_paragraph('Contract Number: CNT-2024-001')
        doc.add_paragraph('')
        
        doc.add_heading('PARTIES:', level=1)
        doc.add_paragraph('• Company A (Buyer): ABC Corporation, 123 Main Street, New York, NY 10001')
        doc.add_paragraph('• Company B (Seller): XYZ Industries, 456 Oak Avenue, Los Angeles, CA 90210')
        doc.add_paragraph('')
        
        doc.add_heading('FINANCIAL TERMS:', level=1)
        doc.add_paragraph('• Total Amount: $50,000.00')
        doc.add_paragraph('• Payment Due Date: February 15, 2024')
        doc.add_paragraph('• Late Fee: 2% per month')
        doc.add_paragraph('')
        
        doc.add_heading('DELIVERABLES:', level=1)
        doc.add_paragraph('• Software Development Services')
        doc.add_paragraph('• Project Timeline: 60 days')
        doc.add_paragraph('• Delivery Date: March 15, 2024')
        doc.add_paragraph('')
        
        doc.add_heading('CONTACT INFORMATION:', level=1)
        doc.add_paragraph('• Buyer Contact: john.doe@abc.com, (555) 123-4567')
        doc.add_paragraph('• Seller Contact: jane.smith@xyz.com, (555) 987-6543')
        doc.add_paragraph('')
        
        doc.add_heading('LEGAL CLAUSES:', level=1)
        doc.add_paragraph('• Termination: Either party may terminate with 30 days written notice')
        doc.add_paragraph('• Liability: Limited to contract amount')
        doc.add_paragraph('• Governing Law: State of New York')
        doc.add_paragraph('')
        
        doc.add_paragraph('This contract is binding upon execution by both parties.')
        doc.add_paragraph('')
        
        doc.add_heading('SIGNATURES:', level=1)
        doc.add_paragraph('_________________    _________________')
        doc.add_paragraph('John Doe             Jane Smith')
        doc.add_paragraph('ABC Corporation      XYZ Industries')
        
        return doc
    
    if __name__ == "__main__":
        doc = create_test_document()
        output_path = "/home/ubuntu/intelligent-document-extraction/test_contract.docx"
        doc.save(output_path)
        print(f"Test document created successfully: {output_path}")
        
except ImportError:
    print("python-docx not available, cannot create test document")
    sys.exit(1)
