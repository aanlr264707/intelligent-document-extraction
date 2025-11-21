#!/usr/bin/env python3
"""
Script to extract text content from the SRS Word document
"""
import sys
import os

try:
    from docx import Document
    
    def extract_text_from_docx(file_path):
        """Extract text from a Word document"""
        doc = Document(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text)
        
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_content.append(" | ".join(row_text))
        
        return "\n".join(text_content)
    
    if __name__ == "__main__":
        srs_path = "/home/ubuntu/attachments/0edcc186-ac77-4a1a-ac4f-dd51fb88b4e9/SRS_Intelligent_Document_Extraction+2.docx"
        
        if os.path.exists(srs_path):
            try:
                content = extract_text_from_docx(srs_path)
                print("=== SRS DOCUMENT CONTENT ===")
                print(content)
                print("\n=== END SRS CONTENT ===")
            except Exception as e:
                print(f"Error extracting content: {e}")
        else:
            print(f"SRS file not found at: {srs_path}")
            
except ImportError:
    print("python-docx not available, cannot extract SRS content")
    sys.exit(1)
